"""Sprint 4 — Advanced Reporting & Analytics.

Batch 1: `GET /api/dashboard/trend/` — one row per day across a window
(the raw series backing the Margin/COGS chart). Batch 2: optional
`?branch_id=` scoping on `dashboard_summary`/`dashboard_top_products`/
`dashboard_trend`, plus the `top_products` product-id fix (grouping by
`product` FK instead of only the free-text `product_name` snapshot) so the
frontend has something to drill down into. Batch 3: `?start_date=&
end_date=&source_document_type=` filtering on `GET
/products/{pk}/cost-movements/`, mirroring `StockMovementFilter`'s shape.
Batch 4: `GET /products/{pk}/cost-movements/export/?export_format=csv|xlsx|pdf`
— the same audit trail as a downloadable file. (Named `export_format`, not
`format` — the latter is reserved by DRF's content negotiation.)

Sprint 5 Batch 6: `GET /reports/recipe-profitability/` and `GET
/reports/ingredient-consumption/` — food-cost/margin reporting over the
RECIPE_CONSUME ledger and the SaleItem.unit_cost snapshot Sprint 5 Batch 5
already writes for recipe-product sales. Pure reads; no new write-side code.
"""

from datetime import date, timedelta
from decimal import Decimal

from django.db import connection
from django.test.utils import CaptureQueriesContext
from django.urls import reverse
from rest_framework import status
from rest_framework.test import APITestCase

from accounts.models import Branch, Tenant, User
from pos.models import (
    Category, InventoryCost, InventoryCostMovement, Product, Sale, SaleItem,
    StockMovement,
)
from pos.services.product_types import ProductType


class _Sprint4ReportingTestBase(APITestCase):
    @classmethod
    def setUpTestData(cls):
        cls.tenant = Tenant.objects.create(name='S4 Reporting Tenant')
        cls.branch_a = Branch.objects.create(tenant=cls.tenant, name='Branch A')
        cls.branch_b = Branch.objects.create(tenant=cls.tenant, name='Branch B')
        cls.manager = User.objects.create_user(
            username='s4mgr', password='pw', role=User.Role.MANAGER, tenant=cls.tenant,
        )
        cls.cashier = User.objects.create_user(
            username='s4cash', password='pw', role=User.Role.CASHIER, tenant=cls.tenant,
        )
        cls.category = Category.objects.create(tenant=cls.tenant, name='Grocery')
        cls.coffee = Product.objects.create(
            tenant=cls.tenant, category=cls.category,
            name='Coffee', barcode='S4-1', sku='SKU-S4-1',
            price=Decimal('700.00'), cost=Decimal('500.00'),
            tax_rate=Decimal('0.00'), stock=Decimal('100'),
        )
        InventoryCost.objects.create(
            tenant=cls.tenant, product=cls.coffee, avg_unit_cost=Decimal('500.0000'),
        )
        cls.tea = Product.objects.create(
            tenant=cls.tenant, category=cls.category,
            name='Tea', barcode='S4-2', sku='SKU-S4-2',
            price=Decimal('200.00'), cost=Decimal('80.00'),
            tax_rate=Decimal('0.00'), stock=Decimal('100'),
        )
        InventoryCost.objects.create(
            tenant=cls.tenant, product=cls.tea, avg_unit_cost=Decimal('80.0000'),
        )

    def setUp(self):
        self.client.force_authenticate(user=self.manager)

    def _make_sale(self, items, *, branch=None, tax_amount=Decimal('0'),
                    sale_status=Sale.Status.COMPLETED, on_date=None):
        """`items`: list of (product, qty, price_each, unit_cost) tuples.
        `on_date`: if given, back-dates `created_at` via a bulk .update()
        (Sale.created_at is auto_now_add, so it can't be set at .create()
        time — the same pattern any Django test needing a specific date
        must use)."""
        subtotal = sum((qty * price_each for _, qty, price_each, _ in items), Decimal('0'))
        total = subtotal + tax_amount
        sale = Sale.objects.create(
            tenant=self.tenant, branch=branch or self.branch_a, cashier=self.manager,
            subtotal=subtotal, tax_amount=tax_amount, total=total,
            method=Sale.Method.CASH, paid=total, change=Decimal('0'), status=sale_status,
        )
        for product, qty, price_each, unit_cost in items:
            SaleItem.objects.create(
                sale=sale, product=product, product_name=product.name,
                barcode=product.barcode or '', qty=qty, price_each=price_each,
                line_total=qty * price_each, unit_cost=unit_cost,
            )
        if on_date:
            Sale.objects.filter(pk=sale.pk).update(
                created_at=timezone_aware(on_date),
            )
            sale.refresh_from_db()
        return sale


def timezone_aware(d):
    from django.utils import timezone
    import datetime
    return timezone.make_aware(datetime.datetime.combine(d, datetime.time(12, 0)))


class DashboardTrendTests(_Sprint4ReportingTestBase):
    def test_per_day_series_across_three_days(self):
        today = date.today()
        d1, d2, d3 = today - timedelta(days=2), today - timedelta(days=1), today

        self._make_sale(
            [(self.coffee, Decimal('10'), Decimal('700.00'), Decimal('500.00'))],
            on_date=d1,
        )
        self._make_sale(
            [(self.tea, Decimal('5'), Decimal('200.00'), Decimal('80.00'))],
            on_date=d2,
        )
        # d3: no sales — must appear as a zero row, not a gap.

        resp = self.client.get(reverse('dashboard-trend'), {
            'start_date': d1.isoformat(), 'end_date': d3.isoformat(),
        })
        self.assertEqual(resp.status_code, status.HTTP_200_OK, resp.content)
        days = {row['date']: row for row in resp.json()['days']}

        self.assertEqual(len(days), 3)
        self.assertEqual(days[d1.isoformat()]['net_revenue'], 7000.0)
        self.assertEqual(days[d1.isoformat()]['cogs'], 5000.0)
        self.assertEqual(days[d1.isoformat()]['gross_profit'], 2000.0)

        self.assertEqual(days[d2.isoformat()]['net_revenue'], 1000.0)
        self.assertEqual(days[d2.isoformat()]['cogs'], 400.0)
        self.assertEqual(days[d2.isoformat()]['gross_profit'], 600.0)

        self.assertEqual(days[d3.isoformat()]['net_revenue'], 0.0)
        self.assertEqual(days[d3.isoformat()]['cogs'], 0.0)
        self.assertEqual(days[d3.isoformat()]['gross_profit'], 0.0)
        self.assertEqual(days[d3.isoformat()]['gross_margin_pct'], 0.0)

    def test_voided_sale_excluded(self):
        today = date.today()
        self._make_sale(
            [(self.tea, Decimal('5'), Decimal('200.00'), Decimal('80.00'))],
            on_date=today, sale_status=Sale.Status.VOIDED,
        )
        resp = self.client.get(reverse('dashboard-trend'), {
            'start_date': today.isoformat(), 'end_date': today.isoformat(),
        })
        day = resp.json()['days'][0]
        self.assertEqual(day['net_revenue'], 0.0)
        self.assertEqual(day['cogs'], 0.0)

    def test_cashier_forbidden(self):
        self.client.force_authenticate(user=self.cashier)
        resp = self.client.get(reverse('dashboard-trend'))
        self.assertEqual(resp.status_code, status.HTTP_403_FORBIDDEN)

    def test_invalid_date_returns_400(self):
        resp = self.client.get(reverse('dashboard-trend'), {'start_date': 'not-a-date'})
        self.assertEqual(resp.status_code, status.HTTP_400_BAD_REQUEST)


class DashboardBranchFilterTests(_Sprint4ReportingTestBase):
    def test_dashboard_summary_branch_filter_narrows_to_one_branch(self):
        self._make_sale(
            [(self.coffee, Decimal('10'), Decimal('700.00'), Decimal('500.00'))],
            branch=self.branch_a,
        )
        self._make_sale(
            [(self.tea, Decimal('5'), Decimal('200.00'), Decimal('80.00'))],
            branch=self.branch_b,
        )

        resp = self.client.get(reverse('dashboard-summary'), {'branch_id': self.branch_a.pk})
        self.assertEqual(resp.status_code, status.HTTP_200_OK, resp.content)
        kpis = resp.json()['kpis']
        self.assertEqual(kpis['net_revenue'], 7000.0)
        self.assertEqual(kpis['cogs'], 5000.0)
        names = {row['name'] for row in resp.json()['top_products']}
        self.assertEqual(names, {'Coffee'})

    def test_no_branch_id_is_unfiltered_regression(self):
        self._make_sale(
            [(self.coffee, Decimal('10'), Decimal('700.00'), Decimal('500.00'))],
            branch=self.branch_a,
        )
        self._make_sale(
            [(self.tea, Decimal('5'), Decimal('200.00'), Decimal('80.00'))],
            branch=self.branch_b,
        )
        resp = self.client.get(reverse('dashboard-summary'))
        kpis = resp.json()['kpis']
        self.assertEqual(kpis['net_revenue'], 8000.0)
        self.assertEqual(kpis['cogs'], 5400.0)

    def test_invalid_branch_id_returns_400(self):
        resp = self.client.get(reverse('dashboard-summary'), {'branch_id': 'nope'})
        self.assertEqual(resp.status_code, status.HTTP_400_BAD_REQUEST)

    def test_unknown_branch_id_returns_404(self):
        resp = self.client.get(reverse('dashboard-summary'), {'branch_id': 999999})
        self.assertEqual(resp.status_code, status.HTTP_404_NOT_FOUND)

    def test_cross_tenant_branch_id_returns_404(self):
        other_tenant = Tenant.objects.create(name='S4 Other Tenant')
        other_branch = Branch.objects.create(tenant=other_tenant, name='Other Branch')
        resp = self.client.get(reverse('dashboard-summary'), {'branch_id': other_branch.pk})
        self.assertEqual(resp.status_code, status.HTTP_404_NOT_FOUND)

    def test_dashboard_trend_branch_filter(self):
        today = date.today()
        self._make_sale(
            [(self.coffee, Decimal('10'), Decimal('700.00'), Decimal('500.00'))],
            branch=self.branch_a, on_date=today,
        )
        self._make_sale(
            [(self.tea, Decimal('5'), Decimal('200.00'), Decimal('80.00'))],
            branch=self.branch_b, on_date=today,
        )
        resp = self.client.get(reverse('dashboard-trend'), {
            'start_date': today.isoformat(), 'end_date': today.isoformat(),
            'branch_id': self.branch_a.pk,
        })
        day = resp.json()['days'][0]
        self.assertEqual(day['net_revenue'], 7000.0)
        self.assertEqual(day['cogs'], 5000.0)


class TopProductsProductIdTests(_Sprint4ReportingTestBase):
    def test_dashboard_summary_top_products_includes_product_id(self):
        self._make_sale([(self.coffee, Decimal('10'), Decimal('700.00'), Decimal('500.00'))])
        resp = self.client.get(reverse('dashboard-summary'))
        row = resp.json()['top_products'][0]
        self.assertEqual(row['id'], self.coffee.pk)
        self.assertEqual(row['name'], 'Coffee')

    def test_standalone_top_products_endpoint_includes_product_id(self):
        today = date.today()
        self._make_sale(
            [(self.coffee, Decimal('10'), Decimal('700.00'), Decimal('500.00'))],
            on_date=today,
        )
        resp = self.client.get(reverse('dashboard-top-products'))
        self.assertEqual(resp.status_code, status.HTTP_200_OK, resp.content)
        row = resp.json()[0]
        self.assertEqual(row['id'], self.coffee.pk)
        self.assertNotIn('product', row, 'internal grouping key should be renamed to id, not duplicated')

    def test_deleted_product_excluded_not_crashed(self):
        """SaleItem.product is SET_NULL — a sale item whose product was
        later deleted must be excluded from the ranking (no id to drill
        into), not crash or appear with id=None."""
        throwaway = Product.objects.create(
            tenant=self.tenant, category=self.category,
            name='Throwaway', barcode='S4-3', sku='SKU-S4-3',
            price=Decimal('50.00'), cost=Decimal('20.00'),
            tax_rate=Decimal('0.00'), stock=Decimal('10'),
        )
        self._make_sale([
            (self.coffee,   Decimal('10'), Decimal('700.00'), Decimal('500.00')),
            (throwaway,     Decimal('1'),  Decimal('50.00'),  Decimal('20.00')),
        ])
        throwaway.delete()

        resp = self.client.get(reverse('dashboard-summary'))
        self.assertEqual(resp.status_code, status.HTTP_200_OK, resp.content)
        names = {row['name'] for row in resp.json()['top_products']}
        self.assertEqual(names, {'Coffee'})


class CostMovementDateFilterTests(_Sprint4ReportingTestBase):
    """Sprint 4 Batch 3 — `?start_date=&end_date=&source_document_type=` on
    `GET /products/{pk}/cost-movements/`. `occurred_at` is auto_now_add,
    so movements are back-dated via a bulk `.update()` after creation, same
    technique `_make_sale`'s `on_date` uses for `Sale.created_at`."""

    def _make_movement(self, *, on_date, source_document_type='purchase_invoice',
                        avg_before=Decimal('500.0000'), avg_after=Decimal('550.0000')):
        inv_cost, _ = InventoryCost.objects.get_or_create(
            tenant=self.tenant, product=self.coffee,
            defaults={'avg_unit_cost': avg_after},
        )
        m = InventoryCostMovement.objects.create(
            tenant=self.tenant, product=self.coffee, inventory_cost=inv_cost,
            quantity_before=Decimal('10'), quantity_received=Decimal('10'),
            unit_cost_received=Decimal('600.00'),
            avg_cost_before=avg_before, avg_cost_after=avg_after,
            source_document_type=source_document_type, source_document_id=1,
        )
        InventoryCostMovement.objects.filter(pk=m.pk).update(
            occurred_at=timezone_aware(on_date),
        )
        m.refresh_from_db()
        return m

    def test_date_range_narrows_results(self):
        today = date.today()
        old, recent = today - timedelta(days=60), today - timedelta(days=1)
        self._make_movement(on_date=old)
        self._make_movement(on_date=recent)

        resp = self.client.get(
            reverse('product-cost-movements', args=[self.coffee.id]),
            {'start_date': (today - timedelta(days=7)).isoformat(), 'end_date': today.isoformat()},
        )
        self.assertEqual(resp.status_code, status.HTTP_200_OK, resp.content)
        self.assertEqual(resp.json()['count'], 1)

    def test_source_document_type_filter(self):
        self._make_movement(on_date=date.today(), source_document_type='purchase_invoice')
        self._make_movement(on_date=date.today(), source_document_type='manual_cost_adjustment')

        resp = self.client.get(
            reverse('product-cost-movements', args=[self.coffee.id]),
            {'source_document_type': 'manual_cost_adjustment'},
        )
        self.assertEqual(resp.json()['count'], 1)
        self.assertEqual(resp.json()['results'][0]['source_document_type'], 'manual_cost_adjustment')

    def test_no_filter_params_is_unfiltered_regression(self):
        self._make_movement(on_date=date.today() - timedelta(days=400))
        self._make_movement(on_date=date.today())
        resp = self.client.get(reverse('product-cost-movements', args=[self.coffee.id]))
        self.assertEqual(resp.json()['count'], 2)

    def test_combined_with_pagination(self):
        for i in range(3):
            self._make_movement(on_date=date.today())
        resp = self.client.get(
            reverse('product-cost-movements', args=[self.coffee.id]),
            {'start_date': date.today().isoformat(), 'page_size': 2},
        )
        self.assertEqual(resp.json()['count'], 3)
        self.assertEqual(len(resp.json()['results']), 2)


class CostHistoryExportTests(_Sprint4ReportingTestBase):
    """Sprint 4 Batch 4 — CSV/Excel/PDF export of the AVCO audit trail."""

    def _make_movement(self, *, on_date=None):
        inv_cost, _ = InventoryCost.objects.get_or_create(
            tenant=self.tenant, product=self.coffee,
            defaults={'avg_unit_cost': Decimal('550.0000')},
        )
        m = InventoryCostMovement.objects.create(
            tenant=self.tenant, product=self.coffee, inventory_cost=inv_cost,
            quantity_before=Decimal('10'), quantity_received=Decimal('10'),
            unit_cost_received=Decimal('600.00'),
            avg_cost_before=Decimal('500.0000'), avg_cost_after=Decimal('550.0000'),
            source_document_type='purchase_invoice', source_document_id=7,
            note='test movement',
        )
        if on_date:
            InventoryCostMovement.objects.filter(pk=m.pk).update(
                occurred_at=timezone_aware(on_date),
            )
        return m

    def test_csv_export_200_correct_headers_and_content(self):
        self._make_movement()
        resp = self.client.get(
            reverse('product-cost-movements-export', args=[self.coffee.id]),
            {'export_format': 'csv'},
        )
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        self.assertIn('text/csv', resp['Content-Type'])
        self.assertIn('attachment; filename=', resp['Content-Disposition'])
        body = resp.content.decode('utf-8-sig')
        self.assertIn('Date,Source,Qty received', body)
        self.assertIn('purchase_invoice #7', body)
        self.assertIn('550.0000', body)

    def test_xlsx_export_200_correct_content_type_and_rows(self):
        self._make_movement()
        resp = self.client.get(
            reverse('product-cost-movements-export', args=[self.coffee.id]),
            {'export_format': 'xlsx'},
        )
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        self.assertIn('spreadsheetml.sheet', resp['Content-Type'])
        self.assertIn('attachment; filename=', resp['Content-Disposition'])

        from io import BytesIO
        from openpyxl import load_workbook
        wb = load_workbook(BytesIO(resp.content))
        ws = wb.active
        self.assertEqual(ws['A1'].value, 'Date')
        self.assertEqual(ws.max_row, 2, 'header row + 1 data row')

    def test_pdf_export_200_correct_content_type(self):
        self._make_movement()
        resp = self.client.get(
            reverse('product-cost-movements-export', args=[self.coffee.id]),
            {'export_format': 'pdf'},
        )
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        self.assertEqual(resp['Content-Type'], 'application/pdf')
        self.assertIn('attachment; filename=', resp['Content-Disposition'])
        self.assertTrue(resp.content.startswith(b'%PDF'))

    def test_empty_result_still_returns_valid_file_not_500(self):
        for fmt, expected_ct in (
            ('csv', 'text/csv'), ('xlsx', 'spreadsheetml.sheet'), ('pdf', 'application/pdf'),
        ):
            resp = self.client.get(
                reverse('product-cost-movements-export', args=[self.coffee.id]),
                {'export_format': fmt},
            )
            self.assertEqual(resp.status_code, status.HTTP_200_OK, f'{fmt} export crashed on empty data')
            self.assertIn(expected_ct, resp['Content-Type'])

    def test_date_range_narrows_export(self):
        old, recent = date.today() - timedelta(days=60), date.today()
        self._make_movement(on_date=old)
        self._make_movement(on_date=recent)
        resp = self.client.get(
            reverse('product-cost-movements-export', args=[self.coffee.id]),
            {'export_format': 'csv', 'start_date': (date.today() - timedelta(days=7)).isoformat(),
             'end_date': date.today().isoformat()},
        )
        body = resp.content.decode('utf-8-sig')
        # Header + exactly 1 data row.
        self.assertEqual(len(body.strip().splitlines()), 2)

    def test_invalid_format_returns_400(self):
        resp = self.client.get(
            reverse('product-cost-movements-export', args=[self.coffee.id]),
            {'export_format': 'txt'},
        )
        self.assertEqual(resp.status_code, status.HTTP_400_BAD_REQUEST)

    def test_cashier_forbidden(self):
        self.client.force_authenticate(user=self.cashier)
        resp = self.client.get(
            reverse('product-cost-movements-export', args=[self.coffee.id]),
            {'export_format': 'csv'},
        )
        self.assertEqual(resp.status_code, status.HTTP_403_FORBIDDEN)


# ── Sprint 5 Batch 6 — Recipe & food-cost reporting ─────────────────────────

class _Batch6ReportingTestBase(_Sprint4ReportingTestBase):
    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        cls.sandwich = Product.objects.create(
            tenant=cls.tenant, category=cls.category,
            name='Chicken Sandwich', barcode='S6-SW', sku='SKU-S6-SW',
            price=Decimal('50.00'), cost=Decimal('0.00'), stock=Decimal('0'),
            product_type=ProductType.RECIPE_PRODUCT,
        )
        cls.burger = Product.objects.create(
            tenant=cls.tenant, category=cls.category,
            name='Beef Burger', barcode='S6-BG', sku='SKU-S6-BG',
            price=Decimal('60.00'), cost=Decimal('0.00'), stock=Decimal('0'),
            product_type=ProductType.RECIPE_PRODUCT,
        )
        cls.chicken = Product.objects.create(
            tenant=cls.tenant, category=cls.category,
            name='Chicken', barcode='S6-CHK', sku='SKU-S6-CHK',
            price=Decimal('0.00'), cost=Decimal('0.00'), stock=Decimal('0'),
        )

    def _make_recipe_movement(self, *, product, branch, qty,
                               avg_unit_cost=Decimal('0.1000'), on_date=None):
        """A `RECIPE_CONSUME` stock movement plus the branch-scoped
        `InventoryCost` row `ingredient_consumption_report` reads to price
        it — mirrors how `SaleSerializer._apply_stock` writes these rows
        (qty stored negative), without going through a real recipe sale."""
        InventoryCost.objects.update_or_create(
            tenant=self.tenant, product=product, branch=branch,
            defaults={'avg_unit_cost': avg_unit_cost},
        )
        mv = StockMovement.objects.create(
            tenant=self.tenant, product=product, branch=branch,
            qty=-qty, movement_type=StockMovement.MovementType.RECIPE_CONSUME,
            source_document_type='sale',
        )
        if on_date:
            StockMovement.objects.filter(pk=mv.pk).update(created_at=timezone_aware(on_date))
            mv.refresh_from_db()
        return mv


class RecipeProfitabilityReportTests(_Batch6ReportingTestBase):
    def test_per_product_aggregate_matches_hand_computed_figures(self):
        self._make_sale([(self.sandwich, Decimal('2'), Decimal('50.00'), Decimal('16.00'))])
        self._make_sale([(self.sandwich, Decimal('1'), Decimal('50.00'), Decimal('16.00'))])
        self._make_sale([(self.burger, Decimal('3'), Decimal('60.00'), Decimal('20.00'))])

        resp = self.client.get(reverse('recipe-profitability'))
        self.assertEqual(resp.status_code, status.HTTP_200_OK, resp.content)
        rows = {r['product_name']: r for r in resp.json()['results']}
        self.assertEqual(set(rows), {'Chicken Sandwich', 'Beef Burger'})

        sw = rows['Chicken Sandwich']
        self.assertEqual(sw['units_sold'], 3.0)
        self.assertEqual(sw['revenue'], 150.0)           # 3 * 50
        self.assertEqual(sw['food_cost'], 48.0)           # 3 * 16
        self.assertEqual(sw['gross_profit'], 102.0)
        self.assertEqual(sw['gross_margin_pct'], 68.0)    # 102/150*100
        self.assertEqual(sw['food_cost_pct'], 32.0)       # 48/150*100

        bg = rows['Beef Burger']
        self.assertEqual(bg['units_sold'], 3.0)
        self.assertEqual(bg['revenue'], 180.0)
        self.assertEqual(bg['food_cost'], 60.0)
        self.assertEqual(bg['gross_profit'], 120.0)
        self.assertEqual(bg['gross_margin_pct'], 66.7)    # round(120/180*100, 1)
        self.assertEqual(bg['food_cost_pct'], 33.3)       # round(60/180*100, 1)

    def test_stock_item_products_excluded(self):
        """Only RECIPE_PRODUCT-typed sale lines appear — a regular stock
        item (Coffee, from the shared Sprint 4 fixture) must never show up
        in a food-cost report."""
        self._make_sale([(self.coffee, Decimal('5'), Decimal('700.00'), Decimal('500.00'))])
        self._make_sale([(self.sandwich, Decimal('1'), Decimal('50.00'), Decimal('16.00'))])
        resp = self.client.get(reverse('recipe-profitability'))
        names = {r['product_name'] for r in resp.json()['results']}
        self.assertEqual(names, {'Chicken Sandwich'})

    def test_branch_filter_narrows(self):
        self._make_sale(
            [(self.sandwich, Decimal('1'), Decimal('50.00'), Decimal('16.00'))], branch=self.branch_a,
        )
        self._make_sale(
            [(self.burger, Decimal('1'), Decimal('60.00'), Decimal('20.00'))], branch=self.branch_b,
        )
        resp = self.client.get(reverse('recipe-profitability'), {'branch_id': self.branch_a.pk})
        self.assertEqual(resp.status_code, status.HTTP_200_OK, resp.content)
        names = {r['product_name'] for r in resp.json()['results']}
        self.assertEqual(names, {'Chicken Sandwich'})

    def test_voided_sale_excluded(self):
        self._make_sale(
            [(self.sandwich, Decimal('1'), Decimal('50.00'), Decimal('16.00'))],
            sale_status=Sale.Status.VOIDED,
        )
        resp = self.client.get(reverse('recipe-profitability'))
        self.assertEqual(resp.json()['results'], [])

    def test_zero_sales_in_window_returns_empty_not_error(self):
        resp = self.client.get(reverse('recipe-profitability'), {
            'start_date': '2020-01-01', 'end_date': '2020-01-02',
        })
        self.assertEqual(resp.status_code, status.HTTP_200_OK, resp.content)
        self.assertEqual(resp.json()['results'], [])

    def test_ordering_param_ascending_margin(self):
        self._make_sale([(self.sandwich, Decimal('1'), Decimal('50.00'), Decimal('16.00'))])  # 68.0%
        self._make_sale([(self.burger, Decimal('1'), Decimal('60.00'), Decimal('55.00'))])     # low margin
        resp = self.client.get(reverse('recipe-profitability'), {'ordering': 'gross_margin_pct'})
        names = [r['product_name'] for r in resp.json()['results']]
        self.assertEqual(names, ['Beef Burger', 'Chicken Sandwich'])

    def test_invalid_ordering_falls_back_to_default(self):
        self._make_sale([(self.sandwich, Decimal('1'), Decimal('50.00'), Decimal('16.00'))])
        resp = self.client.get(reverse('recipe-profitability'), {'ordering': 'not_a_real_field'})
        self.assertEqual(resp.status_code, status.HTTP_200_OK, resp.content)

    def test_cashier_forbidden(self):
        self.client.force_authenticate(user=self.cashier)
        resp = self.client.get(reverse('recipe-profitability'))
        self.assertEqual(resp.status_code, status.HTTP_403_FORBIDDEN)

    def test_invalid_start_date_returns_400(self):
        """Batch 7 verification — `_parse_report_window`'s error branches
        were only ever exercised via `dashboard-summary`/`dashboard-trend`
        in the test suite; this endpoint reuses the same helper but had
        zero coverage of its own error path (measured via `coverage run`,
        not assumed)."""
        resp = self.client.get(reverse('recipe-profitability'), {'start_date': 'not-a-date'})
        self.assertEqual(resp.status_code, status.HTTP_400_BAD_REQUEST, resp.content)

    def test_unknown_branch_id_returns_404(self):
        resp = self.client.get(reverse('recipe-profitability'), {'branch_id': 999999})
        self.assertEqual(resp.status_code, status.HTTP_404_NOT_FOUND)

    def test_non_numeric_branch_id_returns_400(self):
        resp = self.client.get(reverse('recipe-profitability'), {'branch_id': 'not-a-number'})
        self.assertEqual(resp.status_code, status.HTTP_400_BAD_REQUEST, resp.content)

    def test_query_count_does_not_scale_with_sale_row_count(self):
        """Batch 7 verification (B7V-2) — measured: this view aggregates in
        the database (`.values().annotate()`), so 10 sale lines across 2
        products cost a fixed, small number of queries — not one per sale
        line. Pins that shape so a future change (e.g. per-row Python
        post-processing) doesn't silently turn this into an N+1."""
        for _ in range(5):
            self._make_sale([(self.sandwich, Decimal('1'), Decimal('50.00'), Decimal('16.00'))])
            self._make_sale([(self.burger, Decimal('1'), Decimal('60.00'), Decimal('20.00'))])
        with CaptureQueriesContext(connection) as ctx:
            resp = self.client.get(reverse('recipe-profitability'))
        self.assertEqual(resp.status_code, status.HTTP_200_OK, resp.content)
        self.assertEqual(len(resp.json()['results']), 2)
        self.assertLessEqual(
            len(ctx.captured_queries), 5,
            f'recipe-profitability query count scales with row count: '
            f'{len(ctx.captured_queries)} queries for 10 sale lines across 2 products',
        )


class IngredientConsumptionReportTests(_Batch6ReportingTestBase):
    def test_aggregate_matches_hand_computed_figures(self):
        self._make_recipe_movement(
            product=self.chicken, branch=self.branch_a, qty=Decimal('150'),
            avg_unit_cost=Decimal('0.1000'),
        )
        self._make_recipe_movement(
            product=self.chicken, branch=self.branch_a, qty=Decimal('50'),
            avg_unit_cost=Decimal('0.1000'),
        )
        resp = self.client.get(reverse('ingredient-consumption-report'))
        self.assertEqual(resp.status_code, status.HTTP_200_OK, resp.content)
        rows = resp.json()['results']
        self.assertEqual(len(rows), 1)
        row = rows[0]
        self.assertEqual(row['product_name'], 'Chicken')
        self.assertEqual(row['qty_consumed'], 200.0)
        self.assertEqual(row['cost_consumed'], 20.0)  # 200 * 0.10

    def test_branch_scoped_cost_used_per_movement(self):
        """D-09: two branches can carry different average costs for the
        same ingredient — `cost_consumed` must reflect each movement's own
        branch cost, not one global figure."""
        self._make_recipe_movement(
            product=self.chicken, branch=self.branch_a, qty=Decimal('100'),
            avg_unit_cost=Decimal('0.10'),
        )
        self._make_recipe_movement(
            product=self.chicken, branch=self.branch_b, qty=Decimal('100'),
            avg_unit_cost=Decimal('0.20'),
        )
        resp = self.client.get(reverse('ingredient-consumption-report'))
        row = resp.json()['results'][0]
        self.assertEqual(row['qty_consumed'], 200.0)
        self.assertEqual(row['cost_consumed'], 30.0)  # 100*0.10 + 100*0.20

    def test_branch_filter_narrows(self):
        self._make_recipe_movement(
            product=self.chicken, branch=self.branch_a, qty=Decimal('100'),
            avg_unit_cost=Decimal('0.10'),
        )
        self._make_recipe_movement(
            product=self.chicken, branch=self.branch_b, qty=Decimal('100'),
            avg_unit_cost=Decimal('0.20'),
        )
        resp = self.client.get(
            reverse('ingredient-consumption-report'), {'branch_id': self.branch_a.pk},
        )
        self.assertEqual(resp.status_code, status.HTTP_200_OK, resp.content)
        row = resp.json()['results'][0]
        self.assertEqual(row['qty_consumed'], 100.0)
        self.assertEqual(row['cost_consumed'], 10.0)

    def test_non_recipe_consume_movements_excluded(self):
        StockMovement.objects.create(
            tenant=self.tenant, product=self.chicken, branch=self.branch_a,
            qty=Decimal('-5'), movement_type=StockMovement.MovementType.SALE_OUT,
        )
        resp = self.client.get(reverse('ingredient-consumption-report'))
        self.assertEqual(resp.json()['results'], [])

    def test_zero_movements_in_window_returns_empty_not_error(self):
        resp = self.client.get(reverse('ingredient-consumption-report'), {
            'start_date': '2020-01-01', 'end_date': '2020-01-02',
        })
        self.assertEqual(resp.status_code, status.HTTP_200_OK, resp.content)
        self.assertEqual(resp.json()['results'], [])

    def test_ordering_param_ascending_qty(self):
        self._make_recipe_movement(
            product=self.chicken, branch=self.branch_a, qty=Decimal('10'),
            avg_unit_cost=Decimal('1.00'),
        )
        cheese = Product.objects.create(
            tenant=self.tenant, category=self.category,
            name='Cheese', barcode='S6-CHZ', sku='SKU-S6-CHZ',
            price=Decimal('0.00'), cost=Decimal('0.00'), stock=Decimal('0'),
        )
        self._make_recipe_movement(
            product=cheese, branch=self.branch_a, qty=Decimal('40'),
            avg_unit_cost=Decimal('1.00'),
        )
        resp = self.client.get(reverse('ingredient-consumption-report'), {'ordering': 'qty_consumed'})
        names = [r['product_name'] for r in resp.json()['results']]
        self.assertEqual(names, ['Chicken', 'Cheese'])

    def test_cashier_forbidden(self):
        self.client.force_authenticate(user=self.cashier)
        resp = self.client.get(reverse('ingredient-consumption-report'))
        self.assertEqual(resp.status_code, status.HTTP_403_FORBIDDEN)

    def test_invalid_start_date_returns_400(self):
        resp = self.client.get(
            reverse('ingredient-consumption-report'), {'start_date': 'not-a-date'},
        )
        self.assertEqual(resp.status_code, status.HTTP_400_BAD_REQUEST, resp.content)

    def test_unknown_branch_id_returns_404(self):
        resp = self.client.get(reverse('ingredient-consumption-report'), {'branch_id': 999999})
        self.assertEqual(resp.status_code, status.HTTP_404_NOT_FOUND)

    def test_non_numeric_branch_id_returns_400(self):
        resp = self.client.get(
            reverse('ingredient-consumption-report'), {'branch_id': 'not-a-number'},
        )
        self.assertEqual(resp.status_code, status.HTTP_400_BAD_REQUEST, resp.content)

    def test_query_count_does_not_scale_with_movement_row_count(self):
        """Batch 7 verification (B7V-2) — measured: 10 `RECIPE_CONSUME`
        movements sharing one `(product, branch)` pair cost 2 queries (one
        `select_related` fetch for the movements, one cached `InventoryCost`
        lookup), not 10+ — confirming `_unit_cost`'s per-request cache
        actually prevents the N+1 this view's own docstring warns about."""
        for _ in range(10):
            self._make_recipe_movement(
                product=self.chicken, branch=self.branch_a, qty=Decimal('1'),
                avg_unit_cost=Decimal('0.10'),
            )
        with CaptureQueriesContext(connection) as ctx:
            resp = self.client.get(reverse('ingredient-consumption-report'))
        self.assertEqual(resp.status_code, status.HTTP_200_OK, resp.content)
        self.assertEqual(len(resp.json()['results']), 1)
        self.assertLessEqual(
            len(ctx.captured_queries), 5,
            f'ingredient report query count scales with row count: '
            f'{len(ctx.captured_queries)} queries for 10 movements sharing 1 (product,branch)',
        )
